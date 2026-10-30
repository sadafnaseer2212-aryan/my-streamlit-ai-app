import hashlib
import html
import io
import json
from pathlib import Path

import pandas as pd
import streamlit as st
from openai import OpenAI
from pypdf import PdfReader
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

APP_NAME = "AI Document Assistant"
APP_TAGLINE = "Upload. Ask. Understand."
DEFAULT_MODEL = "gpt-4o-mini"
NOT_FOUND = "I couldn't find that in the uploaded documents."

st.set_page_config(page_title=f"{APP_NAME} | {APP_TAGLINE}", page_icon=":material/description:",
                   layout="wide", initial_sidebar_state="expanded")


# ---------------------------------------------------------------------------
# THEME
# Two colour palettes. The toggle in the sidebar only changes which palette is
# written into the CSS variables below — it never touches files, chat or FAISS.
# ---------------------------------------------------------------------------
THEMES = {
    "dark": {"bg": "#18181F", "side": "#1E1E27", "card": "#262631", "accent": "#A78BFA", "text": "#F5F3FF",
             "muted": "#A9A6BC", "line": "#363645", "soft": "#2E2A40", "field": "#1F1F29", "on-accent": "#17131F",
             "user-bg": "#5B47B8", "user-text": "#FFFFFF", "ok": "#5BD49A", "warn": "#F2B45A", "err": "#F2777A"},
    "light": {"bg": "#F7F7FA", "side": "#FFFFFF", "card": "#FFFFFF", "accent": "#8066C2", "text": "#23232E",
              "muted": "#6B6880", "line": "#E6E4EE", "soft": "#F3EFFC", "field": "#FFFFFF", "on-accent": "#FFFFFF",
              "user-bg": "#EDE6FB", "user-text": "#3F2F7A", "ok": "#1E8E5A", "warn": "#A96500", "err": "#C93C3F"},
}

# One stylesheet for the whole app. Every colour comes from a CSS variable (var(--name)).
BASE_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');
.stApp{background:var(--bg);color:var(--text);font-family:'Plus Jakarta Sans',sans-serif}
[data-testid="stHeader"]{background:transparent}
.block-container{max-width:1420px;padding-top:1.4rem;padding-bottom:1.5rem}
.stApp p,.stApp li,.stApp h1,.stApp h2,.stApp h3,.stApp label,[data-testid="stWidgetLabel"] p,[data-testid="stMarkdownContainer"]{color:var(--text)}
.stApp h1,.stApp h2,.stApp h3{font-family:'Plus Jakarta Sans',sans-serif;letter-spacing:-.02em}
[data-testid="stCaptionContainer"],[data-testid="stCaptionContainer"] p{color:var(--muted)!important}
.stApp hr{border-color:var(--line)!important}
/* Sidebar */
[data-testid="stSidebar"]{background:var(--side);border-right:1px solid var(--line)}
[data-testid="stSidebar"] [data-testid="stSidebarHeader"]{height:1rem}
/* Main cards (keyed containers) */
.st-key-upload_card,.st-key-chat_card,.st-key-login_card{background:var(--card);border:1px solid var(--line);border-radius:18px;padding:20px 22px}
/* Buttons */
.stApp button p{color:inherit!important}
.stButton button,.stDownloadButton button,.stFormSubmitButton button{background:var(--card);color:var(--text);border:1px solid var(--line);border-radius:11px;font-weight:600;box-shadow:none}
.stButton button:hover,.stDownloadButton button:hover{border-color:var(--accent);color:var(--accent);background:var(--soft)}
.stApp [data-testid^="stBaseButton-primary"]{background:var(--accent);color:var(--on-accent);border-color:var(--accent)}
.stApp [data-testid^="stBaseButton-primary"]:hover{filter:brightness(1.07);color:var(--on-accent);background:var(--accent)}
.stApp [data-testid^="stBaseButton-primary"] *{color:var(--on-accent)!important}
.st-key-clear_files{display:flex;justify-content:flex-end}
.st-key-clear_files button{white-space:nowrap;min-height:0;padding:2px 10px;border:0;background:transparent;color:var(--muted);font-size:13px}
/* Text inputs, dropdowns and their menus — readable in both themes */
.stApp p,.stApp li,.stApp label,.stApp input,.stApp textarea,.stApp button,[data-testid="stMarkdownContainer"]{font-family:'Plus Jakarta Sans',sans-serif}
[data-testid="stTextInputRootElement"],.stSelectbox [role="group"]{background:var(--field)!important;border:1px solid var(--line)!important;border-radius:10px!important}
[data-testid="stTextInputRootElement"]:focus-within,.stSelectbox [role="group"]:focus-within{border-color:var(--accent)!important}
.stApp input,.stApp textarea{color:var(--text)!important;-webkit-text-fill-color:var(--text)!important;background:transparent!important}
.stApp input::placeholder,.stApp textarea::placeholder{color:var(--muted)!important;-webkit-text-fill-color:var(--muted)!important}
.stSelectbox svg,[data-testid="stTextInputRootElement"] svg{color:var(--muted)!important}
/* Dropdown menu of the model selector (it opens outside the sidebar) */
[data-testid="stSelectboxVirtualDropdown"]{background:var(--card)!important;border:1px solid var(--line)!important;border-radius:10px}
[role="option"]{color:var(--text)!important;background:transparent!important}
[role="option"]:hover,[role="option"][data-focused],[role="option"][aria-selected="true"]{background:var(--soft)!important;color:var(--accent)!important}
/* Toggle and sliders use the accent colour instead of Streamlit red */
.st-key-dark_mode label>div:first-of-type{background:var(--line)!important}
.st-key-dark_mode label[data-selected="true"]>div:first-of-type{background:var(--accent)!important}
[data-testid="stSlider"] [role="group"]>div>div:first-child{background:var(--line)!important}
[data-testid="stSlider"] [role="group"]>div>div[data-rac]{background:var(--accent)!important;box-shadow:none!important}
[data-testid="stSliderThumbValue"] p{color:var(--accent)!important;font-weight:600}
[data-testid="stSliderTickBar"] p{color:var(--muted)!important}
[data-testid="stText"],[data-testid="stText"] *{color:var(--text)!important;font-size:12.5px}
[data-testid="stAppDeployButton"]{display:none}
/* Expanders */
[data-testid="stExpander"] details{background:var(--field);border:1px solid var(--line)!important;border-radius:12px}
[data-testid="stExpander"] summary,[data-testid="stExpander"] summary *{color:var(--text)!important}
[data-testid="stExpander"] summary:hover *{color:var(--accent)!important}
/* File uploader: dashed drop zone; Streamlit's own file list is hidden (we show our own list) */
[data-testid="stFileUploaderDropzone"]{background:var(--soft);border:1.5px dashed var(--accent);border-radius:14px;padding:22px 16px}
[data-testid="stFileUploaderDropzone"] span,[data-testid="stFileUploaderDropzone"] small,[data-testid="stFileUploaderDropzone"] p{color:var(--text)!important}
[data-testid="stFileUploaderDropzoneInstructions"] small{color:var(--muted)!important}
[data-testid="stFileUploaderDropzone"] svg{color:var(--accent)!important;fill:var(--accent)!important}
[data-testid="stFileUploaderDropzone"] button{background:var(--accent)!important;color:var(--on-accent)!important;border:0!important}
[data-testid="stFileUploaderDropzone"] button *{color:var(--on-accent)!important}
[data-testid="stFileChips"],[data-testid="stFileUploaderFile"],[data-testid="stFileUploaderPagination"]{display:none!important}
[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"])::after{content:"＋  Drop more files here, or click to browse";display:block;text-align:center;color:var(--accent);font-weight:600;font-size:13.5px}
/* Chat input */
[data-testid="stChatInput"]{background:var(--field)!important;border:1px solid var(--line)!important;border-radius:14px!important}
[data-testid="stChatInput"]:focus-within{border-color:var(--accent)!important}
[data-testid="stChatInput"]>div{background:transparent!important;border:0!important}
[data-testid="stChatInputSubmitButton"]{background:var(--accent)!important;border-radius:10px!important}
[data-testid="stChatInputSubmitButton"] svg{color:var(--on-accent)!important;fill:var(--on-accent)!important}
/* Chat bubbles: assistant on the left, user on the right */
[data-testid="stChatMessage"]{background:transparent;padding:6px 0;gap:10px}
[data-testid="stChatMessageContent"]{background:var(--soft);border:1px solid var(--line);border-radius:4px 16px 16px 16px;padding:12px 16px;flex:0 1 auto;max-width:86%;margin:0!important}
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]){flex-direction:row-reverse}
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) [data-testid="stChatMessageContent"]{background:var(--user-bg);border-color:var(--user-bg);border-radius:16px 4px 16px 16px}
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) [data-testid="stChatMessageContent"] *{color:var(--user-text)!important}
[data-testid="stChatMessageAvatarUser"]{background:var(--accent)!important;color:var(--on-accent)!important}
[data-testid="stChatMessageAvatarCustom"],[data-testid="stChatMessageAvatarAssistant"]{background:var(--soft)!important;color:var(--accent)!important;border:1px solid var(--line)}
[data-testid="stChatMessageAvatarUser"] *{color:var(--on-accent)!important}
[data-testid="stChatMessageAvatarCustom"] *,[data-testid="stChatMessageAvatarAssistant"] *{color:var(--accent)!important}
/* Small custom pieces written with st.markdown */
.brand{display:flex;gap:11px;align-items:center;margin:4px 0 18px}
.brand-logo{width:40px;height:44px;border-radius:10px;background:var(--accent);color:var(--on-accent);display:grid;place-items:center;font-size:20px;font-weight:700;flex:none}
.brand-name{font-weight:700;font-size:18px;line-height:1.15;color:var(--text);letter-spacing:-.02em}
.brand-sub,.stApp .muted{color:var(--muted);font-size:12.5px}
.card-head{display:flex;gap:12px;align-items:center;margin-bottom:14px}
.card-icon{width:42px;height:42px;border-radius:50%;background:var(--soft);color:var(--accent);display:grid;place-items:center;font-size:20px;flex:none}
.card-title{font-weight:700;font-size:18px;color:var(--text);letter-spacing:-.02em}
.type-strip{display:flex;gap:6px;justify-content:center;flex-wrap:wrap;margin:10px 0 4px}
.tile{display:inline-grid;place-items:center;min-width:36px;height:40px;padding:0 4px;border-radius:7px;color:#fff;font-size:10.5px;font-weight:700;flex:none}
.list-head{font-weight:700;font-size:14.5px;color:var(--text);padding-top:6px}
.file-row{display:flex;gap:12px;align-items:center;border:1px solid var(--line);border-radius:12px;padding:10px 12px;margin:8px 0;background:var(--field)}
.file-name{font-weight:600;font-size:13.5px;color:var(--text);white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.badge{font-size:12px;font-weight:600}
.ok{color:var(--ok)}.warn{color:var(--warn)}.err{color:var(--err)}
.notice{border:1px solid var(--line);border-radius:10px;padding:9px 12px;font-size:13px;font-weight:600;margin:6px 0;background:var(--field)}
.refs{display:flex;flex-wrap:wrap;gap:6px;margin:10px 0 2px}
.ref{font-size:11.5px;border:1px solid var(--line);background:var(--card);color:var(--muted);border-radius:7px;padding:3px 8px}
.welcome{text-align:center;padding:42px 18px 26px}
.welcome-icon{width:60px;height:60px;border-radius:16px;margin:0 auto 12px;background:var(--soft);color:var(--accent);display:grid;place-items:center;font-size:28px}
.welcome-title{font-weight:700;font-size:20px;color:var(--text)}
.welcome-copy{color:var(--muted);font-size:13.5px;max-width:420px;margin:6px auto 0;line-height:1.6}
.status-card{display:flex;gap:10px;align-items:center;border:1px solid var(--line);border-radius:12px;padding:10px 12px;margin:6px 0 10px}
.dot{width:9px;height:9px;border-radius:50%;background:var(--ok);flex:none}
.login-title{font-size:clamp(28px,3vw,40px);font-weight:700;letter-spacing:-.035em;line-height:1.12;color:var(--text);margin:26px 0 10px}
.login-title span{color:var(--accent)}
.app-footer{text-align:center;color:var(--muted);font-size:12px;padding:22px 0 4px}
"""

theme = THEMES["dark" if st.session_state.get("dark_mode", True) else "light"]
css_vars = ";".join(f"--{name}:{value}" for name, value in theme.items())
st.markdown(f"<style>:root{{{css_vars}}}{BASE_CSS}</style>", unsafe_allow_html=True)

# Colour of the small file-type tiles (same in both themes).
TYPE_COLORS = {"PDF": "#E5484D", "DOCX": "#2F6FEB", "XLSX": "#1F9D55", "XLS": "#1F9D55",
               "CSV": "#30A46C", "JSON": "#7C5CE0", "TXT": "#8E8C99"}


def verify_key(key: str) -> tuple[bool, str]:
    """Make a real, inexpensive API call; listing models alone may not verify billing."""
    try:
        client = OpenAI(api_key=key, timeout=20.0, max_retries=0)
        client.chat.completions.create(
            model=DEFAULT_MODEL,
            messages=[{"role": "user", "content": "Reply OK"}],
            max_tokens=5,
        )
        return True, "API key verified successfully."
    except Exception as exc:
        # Never expose an exception containing potentially sensitive details.
        name = type(exc).__name__
        if name == "AuthenticationError":
            return False, "Invalid API key. Please check it and try again."
        if name == "RateLimitError":
            return False, "Rate limit or insufficient API quota. Please review your OpenAI billing and limits."
        if name == "PermissionDeniedError":
            return False, "This API key does not have permission to use the selected model."
        return False, f"Verification failed ({name}). Check your connection and API account."


def rows_to_docs(frame, name, label):
    """Turn a table into Documents of 25 rows each. Every row keeps its column names."""
    docs = []
    frame = frame.fillna("")
    for start in range(0, len(frame), 25):
        batch = frame.iloc[start:start + 25]
        text = "\n".join(f"Row {start + n + 2}: " + "; ".join(f"{col}: {value}" for col, value in row.items())
                         for n, (_, row) in enumerate(batch.iterrows()))
        docs.append(Document(page_content=text, metadata={
            "source": name, "location": f"{label}rows {start + 2}–{start + len(batch) + 1}"}))
    return docs


def docs_from_upload(upload) -> list[Document]:
    name = upload.name
    suffix = Path(name).suffix.lower()
    data = upload.getvalue()
    docs = []
    if suffix == ".pdf":
        # Read PDF pages directly; avoids optional LangChain PDF loader integrations.
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted:
            raise ValueError("Password-protected PDF. Upload an unlocked copy.")
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                docs.append(Document(page_content=page_text, metadata={
                    "source": name, "location": f"Page {page_number}", "pages": len(pdf.pages)}))
    elif suffix == ".docx":
        from docx import Document as WordDocument
        word = WordDocument(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in word.paragraphs if p.text.strip()]
        for table in word.tables:
            paragraphs.extend(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
        docs = [Document(page_content="\n".join(paragraphs), metadata={"source": name, "location": "Document"})]
    elif suffix in (".xlsx", ".xls"):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        for sheet, frame in sheets.items():
            docs.extend(rows_to_docs(frame, name, f"Sheet {sheet}, "))
    elif suffix == ".csv":
        try:
            frame = pd.read_csv(io.BytesIO(data), encoding="utf-8-sig")
        except UnicodeDecodeError:  # Files saved from Excel are often not UTF-8.
            frame = pd.read_csv(io.BytesIO(data), encoding="latin-1")
        docs = rows_to_docs(frame, name, "")
    elif suffix == ".json":
        parsed = json.loads(data.decode("utf-8-sig"))
        if isinstance(parsed, list):
            for start in range(0, len(parsed), 20):
                docs.append(Document(page_content=json.dumps(parsed[start:start + 20], ensure_ascii=False, indent=2),
                                     metadata={"source": name, "location": f"Records {start + 1}–{min(start + 20, len(parsed))}"}))
        else:
            docs = [Document(page_content=json.dumps(parsed, ensure_ascii=False, indent=2),
                             metadata={"source": name, "location": "JSON"})]
    elif suffix == ".txt":
        docs = [Document(page_content=data.decode("utf-8-sig", errors="replace"),
                         metadata={"source": name, "location": "Text"})]
    return [doc for doc in docs if doc.page_content.strip()]


def build_vectorstore(uploads, api_key, chunk_size, chunk_overlap):
    """Read each file, record a status for it, then embed all text into one FAISS index."""
    docs, status = [], {}
    for upload in uploads:
        try:
            file_docs = docs_from_upload(upload)
        except Exception as exc:  # One broken file should not stop the others.
            reason = str(exc) if "Password" in str(exc) else "Could not read (damaged or wrong format)"
            status[upload.name] = ("err", reason)
            continue
        if not file_docs:
            is_pdf = upload.name.lower().endswith(".pdf")
            status[upload.name] = ("warn", "No text found (scanned PDF? needs OCR)" if is_pdf else "No text found")
            continue
        total_pages = file_docs[0].metadata.get("pages")
        if total_pages and len(file_docs) < total_pages:
            status[upload.name] = ("warn", f"Indexed {len(file_docs)} of {total_pages} pages")
        else:
            status[upload.name] = ("ok", "Indexed")
        docs.extend(file_docs)
    if not docs:
        return None, status, 0
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    chunks = splitter.split_documents(docs)
    embeddings = OpenAIEmbeddings(api_key=api_key, model="text-embedding-3-small")
    return FAISS.from_documents(chunks, embeddings), status, len(chunks)


RAG_PROMPT = ChatPromptTemplate.from_messages([
    ("system", "You are AI Document Assistant. Answer ONLY using the document context below.\n"
               f"- If the answer is not in the context, reply exactly: {NOT_FOUND}\n"
               "- For totals, averages, counts or other calculations on spreadsheet/CSV data: only calculate "
               "when the context clearly contains every row needed. Otherwise say that only part of the data "
               "was retrieved, list the values you found, and do not guess a total.\n"
               "- Cite the document name and page, sheet, rows or section for each fact.\n"
               "- Never invent facts or sources.\n\nContext:\n{context}"),
    ("human", "{question}"),
])


def format_docs(docs):
    return "\n\n".join(f"[Source: {d.metadata.get('source', 'Unknown')} | {d.metadata.get('location', 'Unknown')}]\n{d.page_content}" for d in docs)


def show_answer_sources(sources):
    """Source chips under an answer, plus the full excerpts in a collapsed expander."""
    labels = list(dict.fromkeys(s["label"] for s in sources))  # unique, in order
    chips = "".join(f'<span class="ref">▤ {html.escape(label)}</span>' for label in labels)
    st.markdown(f'<div class="refs">{chips}</div>', unsafe_allow_html=True)
    with st.expander(f"View {len(sources)} source excerpts"):
        for source in sources:
            st.markdown(f"**{source['label']}**")
            st.text(source["preview"])


def brand_html():
    return (f'<div class="brand"><div class="brand-logo">▤</div><div><div class="brand-name">{APP_NAME}</div>'
            f'<div class="brand-sub">{APP_TAGLINE}</div></div></div>')


# ---------------------------------------------------------------------------
# SESSION STATE
# ---------------------------------------------------------------------------
if "verified" not in st.session_state:
    st.session_state.verified = False
if "messages" not in st.session_state:
    st.session_state.messages = []
if "uploader_id" not in st.session_state:
    st.session_state.uploader_id = 0  # changing this number empties the file uploader

# ---------------------------------------------------------------------------
# SCREEN 1 — API key verification (separate screen, as before)
# ---------------------------------------------------------------------------
if not st.session_state.verified:
    left, center, right = st.columns([0.8, 1.4, 0.8])
    with center:
        st.markdown(brand_html(), unsafe_allow_html=True)
        st.markdown('<div class="login-title">Chat with your documents.<br><span>Get answers with sources.</span></div>'
                    '<p class="muted" style="font-size:15px;margin-bottom:22px">Upload PDF, Word, Excel, CSV, JSON or TXT files '
                    'and ask questions. Every answer comes from your files.</p>', unsafe_allow_html=True)
        with st.container(key="login_card"):
            st.markdown('<div class="card-title">Connect your OpenAI key</div><p class="muted">We check it with one small '
                        'live request before opening the workspace.</p>', unsafe_allow_html=True)
            with st.form("connect_form", border=False):
                entered_key = st.text_input("OpenAI API key", type="password", placeholder="sk-...",
                                            help="Your key is held in session state, not displayed on the dashboard.")
                submitted = st.form_submit_button("Verify key and open workspace  →", type="primary", use_container_width=True)
            if submitted:
                if not entered_key.strip():
                    st.error("Enter your OpenAI API key to continue.")
                else:
                    with st.spinner("Verifying connection..."):
                        valid, feedback = verify_key(entered_key.strip())
                    if valid:
                        st.session_state.api_key = entered_key.strip()
                        st.session_state.verified = True
                        st.rerun()
                    else:
                        st.error(feedback)
        st.caption("Verification uses a small billable request. Avoid uploading confidential files unless your "
                   "provider and deployment meet your privacy needs.")
    st.markdown(f'<div class="app-footer">{APP_NAME} · {APP_TAGLINE}</div>', unsafe_allow_html=True)
    st.stop()

# ---------------------------------------------------------------------------
# SCREEN 2 — Workspace
# ---------------------------------------------------------------------------
with st.sidebar:
    st.markdown(brand_html(), unsafe_allow_html=True)
    if st.button("New conversation", icon=":material/add_comment:", type="primary", use_container_width=True):
        st.session_state.messages = []
        st.rerun()
    if st.session_state.messages:
        st.download_button("Export chat", icon=":material/download:", use_container_width=True,
                           data=json.dumps(st.session_state.messages, ensure_ascii=False, indent=2),
                           file_name="chat_history.json", mime="application/json")
    st.divider()
    # Theme toggle: its value is read at the top of the script to pick the palette.
    st.toggle("Dark mode", value=True, key="dark_mode")
    model_name = st.selectbox("AI model", ["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini"], key="selected_ai_model")
    with st.expander("Advanced settings", icon=":material/tune:"):
        top_k = st.slider("Relevant sources", 2, 10, 4)
        chunk_size = st.slider("Chunk size", 500, 2000, 1000, step=100)
        chunk_overlap = st.slider("Chunk overlap", 0, 400, 150, step=50)
    st.divider()
    st.markdown('<div class="status-card"><span class="dot"></span><div><b style="font-size:13.5px">OpenAI connected</b>'
                '<div class="muted">Key stays in this session only</div></div></div>', unsafe_allow_html=True)
    if st.button("Disconnect API", icon=":material/logout:", use_container_width=True):
        for item in ("verified", "api_key", "vectorstore", "file_sig", "messages", "file_status", "chunk_count", "index_error"):
            st.session_state.pop(item, None)
        st.rerun()

left, right = st.columns([1, 1.45], gap="medium")

# ----- Left: document upload ------------------------------------------------
with left:
    with st.container(key="upload_card"):
        st.markdown('<div class="card-head"><div class="card-icon">⇪</div><div><div class="card-title">Upload documents</div>'
                    '<div class="muted">PDF, Word, Excel, CSV, JSON or TXT files</div></div></div>', unsafe_allow_html=True)
        uploads = st.file_uploader("Drag and drop your files here", accept_multiple_files=True,
                                   type=["pdf", "docx", "xlsx", "xls", "csv", "json", "txt"],
                                   key=f"uploader_{st.session_state.uploader_id}")
        tiles = "".join(f'<span class="tile" style="background:{color}">{kind}</span>'
                        for kind, color in TYPE_COLORS.items() if kind != "XLS")
        st.markdown(f'<div class="type-strip">{tiles}</div>', unsafe_allow_html=True)

        if uploads:
            # The signature changes only when files or chunk settings change, so switching
            # theme or model never rebuilds the FAISS index.
            signature = (tuple((u.name, hashlib.sha256(u.getvalue()).hexdigest()) for u in uploads), chunk_size,
                         chunk_overlap, hashlib.sha256(st.session_state.api_key.encode()).hexdigest())
            if st.session_state.get("file_sig") != signature:
                with st.spinner("Indexing your documents…"):
                    try:
                        store, status, chunk_count = build_vectorstore(uploads, st.session_state.api_key, chunk_size, chunk_overlap)
                        st.session_state.index_error = None if store else "None of the files contained readable text."
                    except ImportError as exc:
                        store, status, chunk_count = None, {}, 0
                        st.session_state.index_error = f"Missing or incompatible library: {str(exc)[:160]}"
                    except Exception as exc:  # e.g. invalid key or no quota while creating embeddings
                        store, status, chunk_count = None, {}, 0
                        st.session_state.index_error = f"Indexing failed ({type(exc).__name__}). Check API access and quota."
                st.session_state.vectorstore = store
                st.session_state.file_status = status
                st.session_state.chunk_count = chunk_count
                # Saved even on failure, so a failed build is not retried (and billed) on every click.
                st.session_state.file_sig = signature

            col_a, col_b = st.columns([1.5, 1], vertical_alignment="center")
            col_a.markdown(f'<div class="list-head">Uploaded documents ({len(uploads)})</div>', unsafe_allow_html=True)
            with col_b:
                if st.button("Clear all", icon=":material/delete:", key="clear_files"):
                    st.session_state.uploader_id += 1  # a new uploader key starts empty
                    st.rerun()
            for file in uploads:
                kind = Path(file.name).suffix.upper().lstrip(".")
                level, label = st.session_state.file_status.get(file.name, ("err", "Not indexed"))
                icon = {"ok": "✓", "warn": "!", "err": "✕"}[level]
                size = f"{file.size / 1024:,.0f} KB" if file.size < 1024 ** 2 else f"{file.size / 1024 ** 2:,.1f} MB"
                st.markdown(f'<div class="file-row"><span class="tile" style="background:{TYPE_COLORS.get(kind, "#8E8C99")}">{html.escape(kind)}</span>'
                            f'<div style="min-width:0"><div class="file-name">{html.escape(file.name)}</div>'
                            f'<div class="muted">{size} · <span class="badge {level}">{icon} {html.escape(label)}</span></div></div></div>',
                            unsafe_allow_html=True)
            if st.session_state.index_error:
                st.markdown(f'<div class="notice err">✕ {html.escape(st.session_state.index_error)}</div>', unsafe_allow_html=True)
            else:
                ready = sum(1 for level, _ in st.session_state.file_status.values() if level != "err")
                st.markdown(f'<div class="notice ok">✓ Ready · {ready} file(s) · {st.session_state.chunk_count} chunks indexed</div>',
                            unsafe_allow_html=True)
        else:
            # No files: forget the old index (chat history is kept).
            for item in ("vectorstore", "file_sig", "file_status", "chunk_count", "index_error"):
                st.session_state.pop(item, None)
            st.markdown('<p class="muted" style="text-align:center;margin-top:8px">No documents yet. Your files stay in this session only.</p>',
                        unsafe_allow_html=True)

# ----- Right: chat ------------------------------------------------------------
store = st.session_state.get("vectorstore")
with right:
    with st.container(key="chat_card"):
        st.markdown('<div class="card-head"><div class="card-icon">✦</div><div><div class="card-title">Ask about your documents</div>'
                    '<div class="muted">Answers come only from your uploaded files, with sources.</div></div></div>', unsafe_allow_html=True)
        # Messages go in this box; the chat input is placed after it so it stays at the bottom.
        history = st.container(height=560, border=False)
        question = None
        with history:
            if not st.session_state.messages:
                if store:
                    st.markdown('<div class="welcome"><div class="welcome-icon">✦</div><div class="welcome-title">Your documents are ready</div>'
                                '<div class="welcome-copy">Ask a question below, or start with a quick action.</div></div>', unsafe_allow_html=True)
                    a, b = st.columns(2)
                    if a.button("Summarize documents", use_container_width=True):
                        question = "Summarize the main points of the uploaded documents."
                    if b.button("Find key insights", use_container_width=True):
                        question = "What are the key findings in the uploaded documents?"
                    if a.button("Compare sources", use_container_width=True):
                        question = "Compare the main points across the uploaded documents."
                    if b.button("Extract important facts", use_container_width=True):
                        question = "List the important facts with their sources."
                else:
                    st.markdown('<div class="welcome"><div class="welcome-icon">▤</div><div class="welcome-title">Welcome to AI Document Assistant</div>'
                                '<div class="welcome-copy">Upload a PDF, Word file, spreadsheet, CSV, JSON or TXT file on the left. '
                                'Once it is indexed, ask anything about it here.</div></div>', unsafe_allow_html=True)
            for msg in st.session_state.messages:
                with st.chat_message(msg["role"], avatar=None if msg["role"] == "user" else ":material/smart_toy:"):
                    st.markdown(msg["content"])
                    if msg.get("sources"):
                        show_answer_sources(msg["sources"])

        typed = st.chat_input("Ask a question about your documents…" if store else "Upload a document to start chatting",
                              disabled=store is None)
        if typed:
            question = typed

        if question and store:
            st.session_state.messages.append({"role": "user", "content": question})
            with history:
                with st.chat_message("user"):
                    st.markdown(question)
                with st.chat_message("assistant", avatar=":material/smart_toy:"):
                    with st.spinner("Searching your files…"):
                        try:
                            # Retrieve ONCE, and use the same chunks for the answer and the source list.
                            found = store.similarity_search(question, k=top_k)
                            llm = ChatOpenAI(api_key=st.session_state.api_key, model=model_name, temperature=0)
                            chain = RAG_PROMPT | llm | StrOutputParser()
                            answer = chain.invoke({"context": format_docs(found), "question": question})
                            sources = [{"label": f"{d.metadata.get('source', 'Document')} · {d.metadata.get('location', 'Source')}",
                                        "preview": d.page_content[:700]} for d in found]
                            if answer.strip().startswith(NOT_FOUND[:25]):
                                sources = []  # don't show "sources" for an answer that wasn't found
                        except Exception as exc:
                            answer = f"Unable to answer ({type(exc).__name__}). Check API access and quota."
                            sources = []
                    st.markdown(answer)
                    if sources:
                        show_answer_sources(sources)
            st.session_state.messages.append({"role": "assistant", "content": answer, "sources": sources})
            st.rerun()  # redraw so the sidebar "Export chat" button appears right away

st.markdown(f'<div class="app-footer">{APP_NAME} · {APP_TAGLINE} · LangChain × FAISS × OpenAI</div>', unsafe_allow_html=True)
